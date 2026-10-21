"""Tally the author check of paper 11's first-pass codes (19 school notices x 15 items).

Run:  D:\\Research-Publishing\\_tools\\biblio-venv\\Scripts\\python.exe -I tally_checks.py [path\\to\\11-check-sheet.xlsx]
(needs only openpyxl; default path = 11-check-sheet.xlsx next to this script)

Checker's code = first-pass code where Sean typed Y, his correct code where he typed N; rows marked ? or left blank are
excluded and listed. Prints, as the paper reports for the earlier pass ([263]):
  - agreement counting each code as met or not met (Y vs P/N), with Cohen's kappa and Gwet's AC1;
  - agreement on the three levels (Y/P/N), with the differences by item and direction;
  - also lenient (Y/P vs N) agreement and 3-level kappa/AC1, per-item agreement, exact (Clopper-Pearson) 95% bounds;
  - school item rates (strict) and the core-five count under first-pass and checked codes.
This is a non-blind author check of the first-pass coding (the checker saw the first-pass code), not a blind second coder.
"""
import math
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ITEMS = ["D1", "D2", "D3", "D4", "D5", "D6", "D7", "D8a", "D8b", "D9", "K1", "K2a", "K3", "K4", "K5"]


def binom_cdf(k, n, p):
    if k < 0:
        return 0.0
    if k >= n:
        return 1.0
    return sum(math.comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k + 1))


def cp_interval(x, n, a=0.05):
    """Exact Clopper-Pearson 95% interval for x successes out of n."""
    if n == 0:
        return (float("nan"), float("nan"))

    def solve(f):
        lo, hi = 0.0, 1.0
        for _ in range(80):
            mid = (lo + hi) / 2
            if f(mid):
                lo = mid
            else:
                hi = mid
        return (lo + hi) / 2
    lower = 0.0 if x == 0 else solve(lambda p: 1 - binom_cdf(x - 1, n, p) < a / 2)
    upper = 1.0 if x == n else solve(lambda p: binom_cdf(x, n, p) > a / 2)
    return lower, upper


def kappa_ac1(a, b, cats):
    """Cohen's kappa and Gwet's AC1 for two raters over categories cats."""
    n = len(a)
    if n == 0:
        return float("nan"), float("nan"), float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe_k = sum((ca[c] / n) * (cb[c] / n) for c in cats)
    k = (po - pe_k) / (1 - pe_k) if pe_k < 1 else float("nan")
    q = len(cats)
    pi = {c: (ca[c] + cb[c]) / (2 * n) for c in cats}
    pe_g = sum(pi[c] * (1 - pi[c]) for c in cats) / (q - 1)
    ac1 = (po - pe_g) / (1 - pe_g) if pe_g < 1 else float("nan")
    return po, k, ac1


def pct(x):
    return "n/a" if x != x else f"{100 * x:.1f}%"


def f3(x):
    return "n/a (no variation)" if x != x else f"{x:.3f}"


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "11-check-sheet.xlsx"
    ws = load_workbook(path, data_only=True)["Checks"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    c_agree = [h for h in hdr if h and h.startswith("Sean: agree")][0]
    c_corr = [h for h in hdr if h and h.startswith("Sean: correct")][0]
    rows, skipped, problems = [], [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["row"]] is None:
            continue
        unit, item = r[col["unit"]], r[col["item"]]
        fp = str(r[col["first-pass code"]]).strip()[0]
        ag = (str(r[col[c_agree]] or "")).strip().upper()[:1]
        cc = (str(r[col[c_corr]] or "")).strip().upper()[:1]
        if ag == "Y":
            chk = fp
            if cc and cc != fp:
                problems.append(f"row {r[0]} {unit} {item}: agree=Y but a different correct code {cc} was typed (Y taken)")
        elif ag == "N":
            if cc not in ("Y", "P", "N"):
                problems.append(f"row {r[0]} {unit} {item}: N without a correct code (excluded)")
                skipped.append((r[0], unit, item, "N, no code"))
                continue
            if cc == fp:
                problems.append(f"row {r[0]} {unit} {item}: N but the correct code equals the first-pass code {fp} (counted as agreement)")
            chk = cc
        else:
            skipped.append((r[0], unit, item, ag or "blank"))
            continue
        rows.append(dict(row=r[0], unit=unit, item=item, fp=fp, chk=chk, note=r[col["Sean: note"]] or ""))

    total = len(rows) + len(skipped)
    print(f"File: {path}")
    print(f"Rows: {total}; checked (Y or N with code): {len(rows)}; excluded (? / blank / N without code): {len(skipped)}")
    print("Design: non-blind author check of the first-pass codes (checker saw the first-pass code). Not a blind second coder.\n")
    if problems:
        print("Data-entry issues:")
        for p in problems:
            print("  " + p)
        print()
    if not rows:
        print("Nothing checked yet.")
        return

    n = len(rows)
    fp = [r["fp"] for r in rows]
    ck = [r["chk"] for r in rows]
    strict = lambda c: "met" if c == "Y" else "not"
    lenient = lambda c: "not" if c == "N" else "met"

    po2, k2, a2 = kappa_ac1([strict(c) for c in fp], [strict(c) for c in ck], ["met", "not"])
    x2 = sum(strict(a) == strict(b) for a, b in zip(fp, ck))
    lo, hi = cp_interval(x2, n)
    print("MET OR NOT MET (strict: Y vs P/N) - the measure the paper reports first")
    print(f"  agreement {x2}/{n} = {pct(po2)} (exact 95% {pct(lo)} to {pct(hi)}); Cohen's kappa {f3(k2)}; Gwet's AC1 {f3(a2)}")

    po3, k3, a3 = kappa_ac1(fp, ck, ["Y", "P", "N"])
    x3 = sum(a == b for a, b in zip(fp, ck))
    lo, hi = cp_interval(x3, n)
    print("THREE LEVELS (Y / P / N)")
    print(f"  agreement {x3}/{n} = {pct(po3)} (exact 95% {pct(lo)} to {pct(hi)}); Cohen's kappa {f3(k3)}; Gwet's AC1 {f3(a3)}")
    print(f"  differences: {n - x3}, of which {n - x2} change the met-or-not count")

    pol, kl, al = kappa_ac1([lenient(c) for c in fp], [lenient(c) for c in ck], ["met", "not"])
    xl = sum(lenient(a) == lenient(b) for a, b in zip(fp, ck))
    lo, hi = cp_interval(xl, n)
    print("LENIENT (Y/P vs N)")
    print(f"  agreement {xl}/{n} = {pct(pol)} (exact 95% {pct(lo)} to {pct(hi)}); Cohen's kappa {f3(kl)}; Gwet's AC1 {f3(al)}\n")

    print("BY ITEM (n, 3-level agreement, strict agreement, strict kappa, strict AC1, differences first-pass -> checked)")
    by = defaultdict(list)
    for r in rows:
        by[r["item"]].append(r)
    for it in ITEMS:
        rs = by.get(it, [])
        if not rs:
            continue
        a = [r["fp"] for r in rs]
        b = [r["chk"] for r in rs]
        p3 = sum(x == y for x, y in zip(a, b)) / len(rs)
        ps, ks, acs = kappa_ac1([strict(c) for c in a], [strict(c) for c in b], ["met", "not"])
        d = Counter(f"{x}->{y}" for x, y in zip(a, b) if x != y)
        print(f"  {it:4} n={len(rs):2d}  3-level {pct(p3):>6}  strict {pct(ps):>6}  kappa {f3(ks):>18}  AC1 {f3(acs):>6}  {dict(d) if d else ''}")

    dirs = Counter(f"{r['fp']}->{r['chk']}" for r in rows if r["fp"] != r["chk"])
    print(f"\nDIRECTION OF DIFFERENCES (all items): {dict(dirs) if dirs else 'none'}")
    print("  (first pass more generous = Y->P, Y->N, P->N; first pass stricter = N->P, N->Y, P->Y)")
    gen = sum(v for k, v in dirs.items() if k in ("Y->P", "Y->N", "P->N"))
    print(f"  first pass more generous: {gen}; first pass stricter: {sum(dirs.values()) - gen}")

    diffs = [r for r in rows if r["fp"] != r["chk"]]
    if diffs:
        print("\nEACH DIFFERENCE (row, unit, item, first-pass -> checked, note)")
        for r in diffs:
            print(f"  {r['row']:3}  {r['unit']}  {r['item']:4} {r['fp']} -> {r['chk']}  {r['note']}")
    if skipped:
        print("\nEXCLUDED ROWS: " + ", ".join(f"{s[0]} ({s[1]} {s[2]}: {s[3]})" for s in skipped))

    # item rates among the school notices, strict, before and after the check (only units fully checked on that item)
    units = sorted({r["unit"] for r in rows})
    print(f"\nSCHOOL ITEM RATES, STRICT (met = Y), among checked rows (19 school notices in all)")
    for it in ITEMS:
        rs = by.get(it, [])
        if rs:
            print(f"  {it:4}  first pass {sum(r['fp'] == 'Y' for r in rs):2d}/{len(rs)}   checked {sum(r['chk'] == 'Y' for r in rs):2d}/{len(rs)}")
    core = ["D1", "D2", "D3", "D4", "D5"]
    cell = {(r["unit"], r["item"]): r for r in rows}
    full = [u for u in units if all((u, i) in cell for i in core)]
    for lab, key in (("first pass", "fp"), ("checked", "chk")):
        c5 = sum(all(cell[(u, i)][key] == "Y" for i in core) for u in full)
        c5l = sum(all(cell[(u, i)][key] in "YP" for i in core) for u in full)
        mean = sum(sum(cell[(u, i)][key] == "Y" for i in core) for u in full) / len(full) if full else float("nan")
        print(f"  core five (D1-D5 all met), {lab}: strict {c5}/{len(full)}, lenient {c5l}/{len(full)}; mean core items met {mean:.2f}")


if __name__ == "__main__":
    main()
