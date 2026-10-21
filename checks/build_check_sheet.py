"""Build the author check sheet for paper 11 (school privacy notices): all 19 school notices x 15 scored items.

Run:  D:\\Research-Publishing\\_tools\\biblio-venv\\Scripts\\python.exe -I build_check_sheet.py
Reads (read-only) ..\\data\\coding\\pass1_reconciled.csv (first-pass codes, as reconciled, which are the codes the paper's
rates rest on), ..\\data\\coding\\pass1.csv (pre-reconciliation codes, shown for reference), ..\\data\\notices\\<unit>\\notice.txt
and ..\\data\\sample\\sample_NOT_FOR_PUBLICATION.csv (only to check that no institution name or domain leaks into the sheet).
Writes next to this script: 11-check-sheet.xlsx and notice-texts\\<unit>.txt (full text, line-numbered).

No sampling: every school notice (kind = PRIV in strata CB, CI, OT, GV; 19 units) and every scored item used in the paper's
reliability figures (15 items: D1-D9 incl. D8a/D8b, K1, K2a, K3-K5; the unscored D6b and K2b are left out) = 285 rows.
The first pass did not record a quotation per code. The evidence quotes below were picked afterwards (21 Oct 2026) by reading
each notice and noting the line(s) the code most plausibly rests on; for codes of "not met" the closest passage is shown and
labelled as such, or "no matching text found". Quotes are verbatim lines from notice.txt with names, e-mail addresses, URLs
and phone numbers replaced by [NAME] / [EMAIL] / [URL] / [PHONE].
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
NOT = ROOT / "data" / "notices"
OUT_TXT = HERE / "notice-texts"
UNITS = ["CB-010", "CB-016", "CB-025", "CB-033", "CB-038", "CB-046", "CB-051", "CB-053", "CB-058",
         "CI-002", "CI-003", "CI-031", "CI-041", "CI-042", "CI-050", "OT-007", "OT-012", "OT-026", "GV-007"]
ITEMS = ["D1", "D2", "D3", "D4", "D5", "D6", "D7", "D8a", "D8b", "D9", "K1", "K2a", "K3", "K4", "K5"]
LABEL = {"Y": "Y = met", "P": "P = partly met", "N": "N = not met"}

ITEM_TEXT = {
    "D1": ("Itemised description of the personal data (s.5(1)(i); Rule 3(b)(i))",
           "Y: lists at least three distinct, specific data items or categories collected from the user (e.g. name, phone, email, photograph, marks, device ID, location). P: names one or two specific items, or only broad classes ('personal information', 'contact details') without itemising. N: no description, or only 'personal data/information' in general."),
    "D2": ("Specified purposes and the services/uses they enable (s.5(1)(i); Rule 3(b)(ii))",
           "Y: at least two concrete purposes tied to what the service does (e.g. admissions, fee collection, attendance, report cards, personalised lessons). P: only generic purposes ('to provide and improve our services', 'for educational purposes'). N: none."),
    "D3": ("How to withdraw consent (s.6(4); Rule 3(c)(i))",
           "Y: states that consent can be withdrawn AND gives a channel (email, form, in-app setting, link). P: states a right to withdraw/opt out with no channel, or a channel for unsubscribing from marketing only (rule v1.1). N: neither."),
    "D4": ("How to exercise the Data Principal's rights (s.5(1)(ii); Rule 3(c)(ii); Rule 14(1)(a))",
           "Y: names at least access or correction or erasure rights AND a means (email, form, portal, address) to request them. P: lists rights with no means, or a means ('contact us') with no rights named. N: neither."),
    "D5": ("How to complain to the Data Protection Board (s.5(1)(iii); Rule 3(c)(iii))",
           "Y: refers to the Data Protection Board of India (or 'the Board' under the DPDP Act) as a complaint route. P: refers to a regulator/authority generically ('data protection authority') or to a court/cyber-cell only. N: no complaint route to a regulator."),
    "D6": ("Notice available in English or an Eighth Schedule language (s.5(3))",
           "Y: notice is in English or an Eighth Schedule language. (N only if it is in neither.)"),
    "D7": ("Business contact of a DPO or other person who can answer privacy questions (s.8(9); Rule 9)",
           "Y: a named role or designation (DPO, grievance officer, privacy contact) with an email or phone for privacy questions; a dedicated privacy@/grievance@ address also counts as Y (rule v1.1). P: generic 'contact us' address/form with no privacy role (a general support/info address is P). N: none."),
    "D8a": ("Grievance redressal mechanism for privacy grievances described (s.8(10); s.13(1))",
            "Y: a mechanism for privacy grievances is described. P: a mechanism is described but only partly (e.g. scope limited to content/terms, or no clear route). N: none."),
    "D8b": ("Response period stated and 90 days or less (Rule 14(3))",
            "Y: a stated response period of 90 days or less for rights requests, queries or grievances (rule v1.1); a processing time for account deletion alone does not count. N: no such period."),
    "D9": ("Notice understandable on its own (Rule 3(a))",
           "Y: stand-alone privacy page or document that does not depend on another document (terms, admission form, circular) for its core content. P: privacy content is a section inside general terms/website disclaimer. N: only a one-line reference to another document."),
    "K1": ("Notice says how it treats children's / minors' data (s.9(1))",
           "Y: states how it treats data of children/minors/students under 18. N: silent on children (a school notice that never mentions students as children counts as N)."),
    "K2a": ("An age threshold is stated (s.2(f): child = under 18)",
            "Y: an age threshold is stated (any age). N: no age threshold."),
    "K3": ("Parental/guardian consent before processing a child's data (s.9(1); Rule 10(1))",
           "Y: says parent/guardian consent is obtained or required before processing a child's data. P: says parents 'may' review/contact, or consent via admission form without saying so for processing. N: not stated."),
    "K4": ("How the parent is verified as an adult (Rule 10(1))",
           "Y: describes any verification step (ID/age check, OTP to registered adult, DigiLocker token, existing records). P: says it 'verifies' with no method. N: not stated."),
    "K5": ("No tracking, behavioural monitoring or targeted advertising of children (s.9(3))",
           "Y: states it does not track, profile, behaviourally monitor or target ads at children. P: says it limits ads/analytics for children without a flat statement. N: silent, or says it uses ad/behavioural trackers on users generally."),
}

# Evidence per unit and item, in ITEMS order. "a-b,c" = verbatim lines of notice.txt that the code rests on;
# "c:..." = closest passage for a 'not met' code; "-" = no matching text found; "EN" = whole notice in English.
# Optional "|note" is shown after the quote.
EV = {
    "CB-010": "106-113|115-116|c:139|c:91|-|EN|81-84|-|-|87-88|c:93|-|-|-|c:98-99",
    "CB-016": "13-14|16-23|c:31|53,58|-|EN|7|c:58|c:58|6|59-60|61|c:60-61|-|c:29-31,39",
    "CB-025": "15-16,37-38,48-49|33-36|c:39-41|104-108|c:133-136|EN|139-147|139-142|c:133|14-17|-|-|-|-|c:26-28",
    "CB-033": "11-15|28-35|c:58|c:4|-|EN|72-74|-|-|4|66|66|66|66|c:57",
    "CB-038": "54|60|62|73|-|EN|76|-|-|52|-|-|-|-|c:67",
    "CB-046": "46-47|46-49,53-54|-|c:50-52|-|EN|2-3~page header of the website, not in the notice body|-|-|38-41|c:46|-|-|-|c:65-66",
    "CB-051": "6,13|6,13|-|30|-|EN|-|-|-|1-3|-|-|-|-|c:19-20",
    "CB-053": "44-47|53-58|67-69|76|-|EN|75|-|-|41|-|-|-|-|c:64",
    "CB-058": "132-136~the 'examples' list is missing from the extracted text (collapsed content on the page)|139-142|145|148|-|EN|c:1~page header of the website only|-|-|130-131|149-150|150|c:150|-|c:138,147",
    "CI-002": "9-13|18|c:21|22|-|EN|22|-|-|7|-|-|-|-|c:13,15",
    "CI-003": "8-12|15-22|33|c:33-34|-|EN|40,43-44|-|-|3|c:3|-|-|-|c:13-14",
    "CI-031": "57-58|62-68|70|-|-|EN|100,106-109|-|-|41-43|78|78|c:92-95|-|c:74",
    "CI-041": "7-10|13-17|-|31-34|-|EN|-|-|-|4|34|-|34|-|c:10,16",
    "CI-042": "-|42-48|-|-|-|EN|-|-|-|41|c:41|-|-|-|c:47",
    "CI-050": "98-105|107-111|-|122|-|EN|45,122|-|-|93-94|c:96|-|-|-|c:121-122",
    "OT-007": "29-34|45-46|84-85|18,49|c:18|EN|99-106|99-106|-|17-18|73|73|73-74|-|c:45,83",
    "OT-012": "30-41|85-110|340-344|319-336|420-426|EN|412-419|420-423|412-419|4-10|c:34-37|-|-|-|c:116-124",
    "OT-026": "13-24|43-50|35-38|87-93|-|EN|102-106|102-106|-|3-6|c:21-24|-|-|-|-",
    "GV-007": "47-53,55|65-77|c:63|102-105|-|EN|132-133|-|-|13|121-122|122|122-123|c:122|-",
}

# Names to blank out of quotes (institutions, operators, vendors, people, places in names). Matched case-insensitively.
REDACT = ["Smart School Plus", "boscosofttech", "Don Bosco", "Silicon House", "National Public School ITPL", "npsitpl",
          "Zigyasa", "Subhash Gupta", "Aditya Group India", "Aditya", "Brighton International School", "Brighton", "JBS",
          "Pratistha", "EMRS Hirli", "Hirli", "Agnipur", "Nabarangpur", "aakashpublic", "San Academy", "Shiv Jyoti",
          "Trio World school", "TRIO", "Trio Vidyakkendra", "Federal Public School", "Triumph World School", "Educrystal",
          "EduCrystal Education Foundation", "The Cathedral and John Connon School", "Cathedral", "John Connon",
          "Sacred Heart", "St. Augustine’s Day School", "St Augustine Day School", "St. Augustine", "Augustine", "STADS",
          "ALLEN Career Institute Pvt. Ltd.", "ALLEN", "ACIPL", "Dhananjay Srivastava", "The International School Bangalore",
          "International School Bangalore", "TISB", "The British School Society", "The British School", "British School",
          "Satender Pal", "Samagra Shiksha", "Sikkim", "EMMIS", "Tashiling", "Gangtok", "Kolkata", "Bengaluru", "Surya City",
          "Surya Nagar", "Chandapura", "Electronic City", "Dwarka", "Kota"]
RE_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+|\[email\s*protected\]", re.I)
RE_URL = re.compile(r"(https?://\S+|www\.\S+|\b[\w-]+\.(com|in|org|net|edu\.in|co\.in)\b\S*)", re.I)
RE_PHONE = re.compile(r"(\+?\d[\d\s\-/]{7,}\d)")
MAXQ = 900


def redact(s):
    s = RE_EMAIL.sub("[EMAIL]", s)
    s = RE_URL.sub("[URL]", s)
    for t in sorted(REDACT, key=len, reverse=True):
        s = re.sub(r"(?<![A-Za-z])" + re.escape(t) + r"(?![A-Za-z])", "[NAME]", s, flags=re.I)
    s = RE_PHONE.sub(lambda m: "[PHONE]" if sum(c.isdigit() for c in m.group(0)) >= 8 else m.group(0), s)
    return s


def lines_of(spec):
    out = []
    for part in spec.split(","):
        a, _, b = part.partition("-")
        out += list(range(int(a), int(b or a) + 1))
    return out


def evidence(unit, spec, lines):
    note = ""
    if "~" in spec:
        spec, note = spec.split("~", 1)
    if spec == "-":
        return "no matching text found", ""
    if spec == "EN":
        return "(whole notice) The notice text is in English.", "all"
    closest = spec.startswith("c:")
    if closest:
        spec = spec[2:]
    nums = lines_of(spec)
    txt = " ".join(re.sub(r"\s+", " ", lines[n - 1]).strip() for n in nums if lines[n - 1].strip())
    txt = redact(txt)
    if len(txt) > MAXQ:
        txt = txt[:MAXQ].rsplit(" ", 1)[0] + " [...]"
    q = ("Closest text found (not necessarily enough for 'met'): " if closest else "") + "\u201c" + txt + "\u201d"
    if note:
        q += "  [" + note + "]"
    return q, spec


def safe(s):
    s = "" if s is None else str(s)
    return (" " + s) if s[:1] in "=+-@" else s


def main():
    p1 = {r["unit_id"]: r for r in csv.DictReader(open(ROOT / "data" / "coding" / "pass1_reconciled.csv", encoding="utf-8-sig"))}
    p1raw = {r["unit_id"]: r for r in csv.DictReader(open(ROOT / "data" / "coding" / "pass1.csv", encoding="utf-8-sig"))}
    school = [u for u, r in p1.items() if r["kind"] == "PRIV" and u[:2] in ("CB", "CI", "OT", "GV")]
    if sorted(school) != sorted(UNITS):
        sys.exit(f"school notice list differs from UNITS: {sorted(school)}")
    OUT_TXT.mkdir(exist_ok=True)
    rows = []
    for u in UNITS:
        raw = (NOT / u / "notice.txt").read_text(encoding="utf-8", errors="replace")
        lines = raw.split("\n")
        with open(OUT_TXT / f"{u}.txt", "w", encoding="utf-8", newline="\r\n") as f:
            f.write(f"Unit {u} - full extracted text of the privacy notice page (from data\\notices\\{u}\\notice.txt), {len(raw.split())} words.\n")
            f.write("Line numbers on the left match the 'lines' column of the check sheet. Page menus and footers are included as extracted.\n")
            f.write("FOR RESEARCH ONLY: this text names the institution. Do not share or upload it.\n")
            f.write("-" * 100 + "\n")
            for i, ln in enumerate(lines, 1):
                f.write(f"{i:4d} | {ln}\n")
        specs = EV[u].split("|")
        assert len(specs) == len(ITEMS), u
        for it, sp in zip(ITEMS, specs):
            code = p1[u][it]
            q, where = evidence(u, sp, lines)
            rows.append(dict(unit=u, item=it, code=code, raw=p1raw[u][it], q=q, where=where))
    # leak check: no institution name tokens or domains from the sample file in the sheet text
    sample = {r["unit_id"]: r for r in csv.DictReader(open(ROOT / "data" / "sample" / "sample_NOT_FOR_PUBLICATION.csv", encoding="utf-8-sig"))}
    common = {"school", "public", "international", "senior", "secondary", "academy", "convent", "english", "medium", "model",
              "vidyalaya", "world", "national", "higher", "group", "india", "indian", "central", "education", "educational",
              "society", "college", "institute", "the", "and", "day", "residential", "high", "girls", "boys", "co-ed", "trust",
              "career", "department", "government", "foundation", "children", "school,", "heart", "federal"}  # "federal": FTC in a CI-031 quote; CI-003 name is redacted
    alltext = " ".join(r["q"] for r in rows).lower()
    leaks = []
    for u in UNITS:
        s = sample.get(u, {})
        toks = {t.strip(".,()'\"").lower() for t in (s.get("name", "") or "").split()}
        for t in toks:
            if len(t) >= 4 and t not in common and re.search(r"(?<![a-z])" + re.escape(t) + r"(?![a-z])", alltext):
                leaks.append((u, t))
        dom = re.sub(r"^https?://(www\.)?", "", s.get("url", "") or "").split("/")[0].lower()
        if dom and dom in alltext:
            leaks.append((u, dom))
    if leaks:
        sys.exit(f"possible name/URL leak in quotes: {leaks}")

    wb = Workbook()
    yellow = PatternFill("solid", fgColor="FFF2A8")
    grey = PatternFill("solid", fgColor="E7E6E6")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "How to"
    howto = [
        "Paper 11 (school privacy notices): author check of the first-pass codes on all 19 school notices",
        "",
        "What this is: each row is one school notice x one checklist item. The first-pass code (Y = met, P = partly met, N = not met) is already filled in, with the passage of the notice it rests on. You say whether you agree.",
        "This is a non-blind check by the author of the first-pass coding (you can see the first-pass code). It is not an independent, blind second coder, and the paper must describe it that way.",
        "",
        "For each row on the 'Checks' sheet:",
        "1. Read the item and its coding rule (columns D-E).",
        "2. Read the evidence quote (column G). If you need more context, open notice-texts\\<unit>.txt and go to the line numbers in column H. The rule is about what the notice SAYS, not what the school does.",
        "3. Column J: type Y if the first-pass code (column F) is right under the rule, N if it is wrong, ? if you cannot decide.",
        "4. If you typed N: in column K put the correct code (Y, P or N).",
        "5. Column L (optional): a short note, e.g. which line shows the code is wrong, or why you are unsure.",
        "",
        "Tips: rows are grouped by notice (15 rows each), so read one notice once and then do its 15 rows. A 'Closest text found' quote on an N row is shown only to help you judge; it may not meet the rule. 'no matching text found' means nothing relevant was found when the evidence was picked; if you find something in the full text, mark N and give the line.",
        "The evidence quotes were picked after the first pass (the first pass did not record them), so a weak quote is not by itself a reason to say N: judge the code against the whole notice.",
        "Names, e-mails, web addresses and phone numbers are blanked in this sheet as [NAME], [EMAIL], [URL], [PHONE]. The full texts in notice-texts\\ are not blanked: keep them for research only, do not share or upload them.",
        "",
        "Time: about 2.5 to 3.5 hours in all (19 notices, 285 rows; the notices total about 25,000 words). It can be done in several sittings; save as you go.",
        "When done: save this file under the same name and tell Claude: 11 check done.",
        "",
        "Sheets: Checks (the 285 rows), Items (the 15 items and full rules), Notices (the 19 notices with word counts and the first-pass note).",
        "Codes used: the reconciled first-pass codes (pass1_reconciled.csv), which are the codes behind the paper's rates. Column I shows the code before the one rule tightening (pass1.csv) where it differs.",
    ]
    for i, t in enumerate(howto, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 130

    ws = wb.create_sheet("Checks")
    hdr = ["row", "unit", "item", "item (what the law asks)", "coding rule", "first-pass code", "evidence quote from the notice",
           "lines in notice-texts\\<unit>.txt", "code before rule tightening (if different)",
           "Sean: agree? (Y / N / ?)", "Sean: correct code if N (Y / P / N)", "Sean: note"]
    ws.append(hdr)
    for i, r in enumerate(rows, 1):
        it = ITEM_TEXT[r["item"]]
        ws.append([i, r["unit"], r["item"], it[0], it[1], LABEL[r["code"]], safe(r["q"]),
                   (f"{r['unit']}.txt, lines {r['where']}" if r["where"] not in ("", "all") else ("whole notice" if r["where"] == "all" else "")),
                   (LABEL.get(r["raw"], r["raw"]) if r["raw"] != r["code"] else ""), None, None, None])
    widths = [6, 9, 7, 30, 48, 14, 80, 18, 14, 12, 14, 30]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.alignment = wrap
    for c in ws[1]:
        c.font = bold
        c.fill = grey
    for rr in range(1, ws.max_row + 1):
        for col in ("J", "K", "L"):
            ws[f"{col}{rr}"].fill = yellow
    ws.freeze_panes = "D2"
    n = ws.max_row
    dv1 = DataValidation(type="list", formula1='"Y,N,?"', allow_blank=True, showErrorMessage=True,
                         errorTitle="Y, N or ?", error="Type Y (agree), N (disagree) or ? (unsure).")
    dv2 = DataValidation(type="list", formula1='"Y,P,N"', allow_blank=True, showErrorMessage=True,
                         errorTitle="Y, P or N", error="Correct code: Y (met), P (partly met) or N (not met).")
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"J2:J{n}")
    dv2.add(f"K2:K{n}")
    ws.auto_filter.ref = f"A1:L{n}"

    ws = wb.create_sheet("Items")
    ws.append(["item", "what the law asks", "coding rule (Y / P / N)"])
    for it in ITEMS:
        ws.append([it, ITEM_TEXT[it][0], ITEM_TEXT[it][1]])
    for j, w in enumerate([7, 45, 110], 1):
        ws.column_dimensions[chr(64 + j)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = wrap
    for c in ws[1]:
        c.font = bold
        c.fill = grey
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Notices")
    ws.append(["unit", "stratum", "words in extracted page", "full text file", "first-pass note (no names)"])
    strat = {"CB": "CBSE school", "CI": "CISCE school", "OT": "other-board school", "GV": "government school system"}
    for u in UNITS:
        nw = len((NOT / u / "notice.txt").read_text(encoding="utf-8", errors="replace").split())
        ws.append([u, strat[u[:2]], nw, f"notice-texts\\{u}.txt", safe(redact(p1[u]["note"]))])
    for j, w in enumerate([9, 22, 12, 24, 100], 1):
        ws.column_dimensions[chr(64 + j)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = wrap
    for c in ws[1]:
        c.font = bold
        c.fill = grey
    ws.freeze_panes = "A2"

    out = HERE / "11-check-sheet.xlsx"
    wb.save(out)
    from collections import Counter
    print(f"wrote {out} with {len(rows)} rows; codes {Counter(r['code'] for r in rows)}; "
          f"no-match quotes {sum(r['q'].startswith('no matching') for r in rows)}; closest {sum(r['q'].startswith('Closest') for r in rows)}")


if __name__ == "__main__":
    main()
