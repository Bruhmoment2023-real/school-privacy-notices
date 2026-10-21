"""Apply the filled check sheet: public copy of the sheet, final codes and change log.

Run:  python -I checks/apply_check.py <filled sheet with notice quotes, kept for research only> <repo>
Writes
  checks/11-check-sheet.filled.xlsx      public copy: evidence quotes removed, quoted notice wording in the
                                         checker's notes replaced by [quote removed] (quotes of the coding rules, listed in KEEP, are kept)
  data/coding/codes_final_checked.csv    pass1_reconciled.csv with every checked correction applied
  data/coding/check_changes_log.csv      one row per change (unit, item, first pass, checked, checker's note, public form)
pass1_reconciled.csv itself is not changed.
"""
import csv
import io
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

src, repo = Path(sys.argv[1]), Path(sys.argv[2])
REMOVED = "[removed: notice text is kept for research only]"
QUOTE = re.compile(r"(?:(?<=^)|(?<=[\s(\[:;,=/]))'((?:[^'\n]|(?<=[A-Za-z])'(?=[A-Za-z]))+?)'(?=[\s).,;:\]/]|$)")
# quotations of the coding rules themselves (not of a notice) are kept
KEEP = {"portal", "states a right to withdraw/opt out", "states a right to withdraw/opt out with no channel",
        "a means ... with no rights named"}

wb = load_workbook(src)


def strip(note):
    if not note:
        return note
    return QUOTE.sub(lambda m: m.group(0) if m.group(1) in KEEP else "[quote removed]", str(note))


ws = wb["Checks"]
ws.cell(1, 7).value = "evidence quote from the notice (removed in the public copy)"
fix = {}
for r in range(2, ws.max_row + 1):
    unit, item = ws.cell(r, 2).value, ws.cell(r, 3).value
    if not unit:
        continue
    ws.cell(r, 7).value = REMOVED
    note = ws.cell(r, 12).value
    ws.cell(r, 12).value = strip(note)
    yn = str(ws.cell(r, 10).value or "").strip().upper()
    corr = str(ws.cell(r, 11).value or "").strip().upper()
    first = str(ws.cell(r, 6).value or "").strip()[:1].upper()
    if yn == "N" and corr in ("Y", "P", "N") and corr != first:
        fix[(unit, item)] = (first, corr, " ".join(str(ws.cell(r, 12).value or "").split()))
wb.save(repo / "checks" / "11-check-sheet.filled.xlsx")

raw = (repo / "data/coding/pass1_reconciled.csv").read_bytes().decode("utf-8")
lines = raw.splitlines(keepends=True)          # untouched lines are copied as they are
head = next(csv.reader([lines[0]]))
log, out = [], [lines[0]]
for line in lines[1:]:
    body = line.rstrip(chr(13) + chr(10))
    row = next(csv.reader([body]))
    changed = False
    for i, item in enumerate(head):
        k = (row[0], item)
        if k in fix:
            first, corr, note = fix.pop(k)
            assert row[i] == first, (k, row[i], first)
            row[i] = corr
            changed = True
            log.append([row[0], item, first, corr, note + (" [note cut at 500 characters in the sheet]" if len(note) >= 495 else "")])
    if changed:
        buf = io.StringIO()
        csv.writer(buf, lineterminator="").writerow(row)
        line = buf.getvalue() + line[len(body):]
    out.append(line)
assert not fix, f"corrections with no matching code: {fix}"
(repo / "data/coding/codes_final_checked.csv").write_bytes("".join(out).replace(chr(13) + chr(10), chr(10)).encode("utf-8"))  # LF, as committed
with open(repo / "data/coding/check_changes_log.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    w.writerow(["unit_id", "item", "first_pass", "checked", "note"])
    w.writerows(log)
print(f"{len(log)} codes changed")
for x in log:
    print(" ", x[0], x[1], x[2], "->", x[3])
